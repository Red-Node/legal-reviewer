"""DOCX integration, archive validation, source-location and browser-memory tests."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from fastapi import UploadFile
from fastapi.testclient import TestClient
import pymupdf

from main import app, analyze_document, MAX_FILE_SIZE
from docx_reader import DocxError, extract_docx
from quote_verifier import verify_docx_evidence
import test_review


def docx_bytes(document=None):
    if document is None:
        document = Document()
        document.add_heading('Section 1.1 — Payment Terms', 1)
        document.add_paragraph('"Business Day" means Monday through Friday.')
        document.add_paragraph('Use the Delivery Window pursuant to Section 9.9.')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Pay on the Business Day.'
        table.cell(0, 1).text = 'See Section 1.1.'
        document.add_paragraph('Section 2.1 — Delivery Terms')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def alter_package(contents, changes=None, additions=None):
    result = BytesIO()
    with ZipFile(BytesIO(contents)) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item)
            if changes and item.filename in changes:
                value = changes[item.filename](value)
            target.writestr(item.filename, value)
        for name, value in (additions or {}).items():
            target.writestr(name, value)
    return result.getvalue()


class DocxTests(unittest.TestCase):
    def setUp(self):
        for target in ('socket.socket.connect', 'ai_reviewer.Anthropic'):
            guard = patch(target, side_effect=AssertionError('DOCX tests must not contact an API.'))
            guard.start()
            self.addCleanup(guard.stop)
        self.client = TestClient(app)
        self.contents = docx_bytes()

    def upload(self, contents=None, filename='sample.docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'):
        return self.client.post('/analyze', files={'file': (filename, self.contents if contents is None else contents, mime)})

    def test_valid_upload_no_fake_pages(self):
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['document_type'], 'docx')
        self.assertEqual(data['filename'], 'sample.docx')
        self.assertFalse(data['ai_supported'])
        self.assertNotIn('total_pages', data)
        self.assertNotIn('pages', data)
        def check(value):
            if isinstance(value, dict):
                self.assertNotIn('page', value)
                for item in value.values(): check(item)
            if isinstance(value, list):
                for item in value: check(item)
        check(data)
        self.assertEqual(response.headers['cache-control'], 'no-store')

    def test_order_paragraph_table_paragraph(self):
        sources = extract_docx(self.contents)
        self.assertEqual([s['location_id'] for s in sources], ['docx-p1', 'docx-p2', 'docx-p3', 'docx-t1-r1-c1-p1', 'docx-t1-r1-c2-p1', 'docx-p4'])
        self.assertEqual(sources[3]['text'], 'Pay on the Business Day.')
        self.assertEqual(sources[4]['location_label'], 'Table 1, row 1, cell 2, Paragraph 1')

    def test_nested_table_order(self):
        document = Document()
        document.add_paragraph('Before')
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = 'Cell before'
        cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Nested'
        cell.add_paragraph('Cell after')
        document.add_paragraph('After')
        self.assertEqual([s['text'] for s in extract_docx(docx_bytes(document)) if s['text']], ['Before', 'Cell before', 'Nested', 'Cell after', 'After'])

    def test_merged_cells_not_duplicated(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 1)).text = 'Merged source'
        self.assertEqual([s['text'] for s in extract_docx(docx_bytes(document))], ['Merged source'])

    def test_empty_paragraphs_have_stable_locations(self):
        document = Document()
        document.add_paragraph('')
        document.add_paragraph('Readable')
        self.assertEqual(extract_docx(docx_bytes(document))[1]['location_id'], 'docx-p2')

    def test_literal_headings_and_references(self):
        data = self.upload().json()
        self.assertEqual([s['number'] for s in data['sections']], ['1.1', '2.1'])
        self.assertEqual([r['number'] for r in data['broken_references']], ['9.9'])
        self.assertEqual(data['broken_references'][0]['location_label'], 'Paragraph 3')
        self.assertEqual(data['references'][1]['location_id'], 'docx-t1-r1-c2-p1')

    def test_defined_and_undefined_terms(self):
        data = self.upload().json()
        self.assertEqual([t['term'] for t in data['defined_terms']], ['Business Day'])
        self.assertEqual([t['term'] for t in data['potentially_undefined_terms']], ['Delivery Window'])
        self.assertEqual(data['potentially_undefined_terms'][0]['location_id'], 'docx-p3')

    def test_later_definitions_and_sections(self):
        document = Document()
        document.add_paragraph('Use the Business Day pursuant to Section 1.1.')
        document.add_paragraph('Section 1.1 — Payment Terms')
        document.add_paragraph('“Business Day” shall mean Monday through Friday.')
        data = self.upload(docx_bytes(document)).json()
        self.assertEqual(data['broken_references'], [])
        self.assertEqual(data['potentially_undefined_terms'], [])

    def test_no_definitions_warning(self):
        document = Document()
        document.add_paragraph('Use the Delivery Window.')
        data = self.upload(docx_bytes(document)).json()
        self.assertTrue(data['term_validation_warning'])
        self.assertEqual(data['potentially_undefined_terms'], [])

    def test_hyperlink_text_preserved_without_fetching(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE
        document = Document()
        paragraph = document.add_paragraph('See ')
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to('https://example.invalid', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run, text = OxmlElement('w:r'), OxmlElement('w:t')
        text.text = 'Section 9.9'
        run.append(text); link.append(run); paragraph._p.append(link)
        self.assertEqual(extract_docx(docx_bytes(document))[0]['text'], 'See Section 9.9')

    def test_uppercase_extension_and_unreliable_mime(self):
        self.assertEqual(self.upload(filename='SAMPLE.DOCX', mime='application/octet-stream').status_code, 200)

    def test_invalid_docx(self):
        for contents in (b'ordinary text', b'PK\x03\x04not really a ZIP', b''):
            with self.subTest(contents=contents):
                self.assertEqual(self.upload(contents).status_code, 400)

    def test_zip_not_word(self):
        data = BytesIO()
        with ZipFile(data, 'w') as archive:
            archive.writestr('not-a-document.txt', 'text')
        self.assertEqual(self.upload(data.getvalue()).status_code, 400)

    def test_legacy_doc_clear_error(self):
        response = self.upload(filename='contract.doc')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Legacy .doc', response.json()['detail'])

    def test_other_extensions(self):
        self.assertEqual(self.upload(filename='contract.docm').status_code, 400)

    def test_pdf_renamed_docx_and_docx_renamed_pdf(self):
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), 'Sample PDF')
            self.assertEqual(self.upload(pdf.tobytes()).status_code, 400)
        self.assertEqual(self.upload(filename='fake.pdf', mime='application/pdf').status_code, 400)

    def test_empty_body_rejected(self):
        self.assertEqual(self.upload(docx_bytes(Document())).status_code, 400)

    def test_invalid_xml_safe_message(self):
        contents = alter_package(self.contents, {'word/document.xml': lambda _: b'<invalid>PRIVATE CONTENT'})
        response = self.upload(contents)
        self.assertEqual(response.status_code, 400)
        self.assertNotIn('PRIVATE', response.text)

    def test_macros_rejected(self):
        contents = alter_package(self.contents, additions={'word/vbaProject.bin': b'fake macro'})
        self.assertEqual(self.upload(contents).status_code, 400)

    def test_macro_content_type_rejected(self):
        contents = alter_package(self.contents, {'[Content_Types].xml': lambda b: b.replace(b'wordprocessingml.document.main+xml', b'wordprocessingml.template.main+xml')})
        self.assertEqual(self.upload(contents).status_code, 400)

    def test_xml_entity_declarations_rejected(self):
        for declarations in (b'<!DOCTYPE document [<!ENTITY x "secret">]>', b'<!ENTITY x "secret">'):
            contents = alter_package(self.contents, {'word/document.xml': lambda b: declarations + b})
            self.assertEqual(self.upload(contents).status_code, 400)

    def test_expanded_zip_limit(self):
        with patch('docx_reader.MAX_EXPANDED_BYTES', 10):
            response = self.upload()
        self.assertEqual(response.status_code, 400)
        self.assertIn('safe processing limit', response.json()['detail'])

    def test_archive_member_limit(self):
        with patch('docx_reader.MAX_ARCHIVE_MEMBERS', 2):
            self.assertEqual(self.upload().status_code, 400)

    def test_upload_size_limit(self):
        response = self.upload(b'x' * (MAX_FILE_SIZE + 1))
        self.assertEqual(response.status_code, 413)
        self.assertIn('10 MB', response.json()['detail'])

    def test_upload_closed_even_on_failure(self):
        for contents in (self.contents, b'broken'):
            stream = BytesIO(contents)
            upload = UploadFile(filename='sample.docx', file=stream)
            try:
                analyze_document(upload)
            except Exception:
                pass
            self.assertTrue(stream.closed)

    def test_deterministic_locations(self):
        self.assertEqual(extract_docx(self.contents), extract_docx(self.contents))

    def test_pdf_response_unchanged(self):
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), 'Section 1.1 Payment Terms')
            data = self.upload(pdf.tobytes(), filename='sample.pdf', mime='application/pdf').json()
        self.assertEqual(data['total_pages'], 1)
        self.assertEqual(data['sections'][0]['page'], 1)
        self.assertNotIn('locations', data)

    def test_docx_locations_rejected_by_demo_endpoint(self):
        response = self.client.post('/ai-review-demo', json={'locations': extract_docx(self.contents)})
        self.assertEqual(response.status_code, 400)

    def test_docx_locations_not_accepted_as_pdf_pages(self):
        response = self.client.post('/ai-review-demo', json={'pages': extract_docx(self.contents)})
        self.assertEqual(response.status_code, 400)

    def test_docx_real_ai_rejected_even_when_local_gate_enabled(self):
        with patch('main._require_local_development'), patch('main.get_real_review_settings', return_value={'enabled': True}), patch('main.review_contract_real') as review:
            response = self.client.post('/ai-review-real', json={'pages': extract_docx(self.contents), 'confirmed': True}, headers={'origin': 'http://testserver'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('not supported for DOCX', response.json()['detail'])
        review.assert_not_called()

    def test_public_real_ai_gate_unchanged_for_docx(self):
        with patch('main.is_public_demo_only', return_value=True):
            response = self.client.post('/ai-review-real', json={'locations': extract_docx(self.contents), 'confirmed': True})
        self.assertEqual(response.status_code, 403)

    def test_correct_docx_evidence(self):
        sources = extract_docx(self.contents)
        evidence = [{'quote': 'See Section 1.1.', 'location_id': sources[4]['location_id']}]
        result = verify_docx_evidence(evidence, sources)
        self.assertEqual(result[0]['verification_status'], 'verified')
        self.assertEqual(result[0]['match_count'], 1)
        self.assertNotIn('page', result[0])
        self.assertEqual(evidence, [{'quote': 'See Section 1.1.', 'location_id': sources[4]['location_id']}])

    def test_wrong_docx_location_and_invented_quote(self):
        sources = extract_docx(self.contents)
        for quote, identifier in [('See Section 1.1.', 'docx-p1'), ('Invented.', 'docx-p1'), ('Text', 'nonexistent')]:
            result = verify_docx_evidence([{'quote': quote, 'location_id': identifier}], sources)
            self.assertEqual(result[0]['verification_status'], 'verification_failed')
            self.assertTrue(result[0]['failure_reason'])

    def test_page_numbers_cannot_verify_docx_evidence(self):
        result = verify_docx_evidence([{'quote': 'See Section 1.1.', 'location_id': 'docx-t1-r1-c2-p1', 'page': 1}], extract_docx(self.contents))
        self.assertEqual(result[0]['verification_status'], 'verification_failed')

    def test_whitespace_repeated_and_empty_docx_quotes(self):
        sources = [{'location_id': 'docx-p1', 'text': 'Notice\n  promptly. Notice promptly.'}]
        result = verify_docx_evidence([{'quote': 'Notice promptly.', 'location_id': 'docx-p1'}, {'quote': '', 'location_id': 'docx-p1'}], sources)
        self.assertEqual(result[0]['match_count'], 2)
        self.assertEqual(result[1]['verification_status'], 'verification_failed')

    def test_docx_evidence_does_not_change_punctuation(self):
        result = verify_docx_evidence([{'quote': 'Pay 10 days!', 'location_id': 'docx-p1'}], [{'location_id': 'docx-p1', 'text': 'Pay 10 days.'}])
        self.assertEqual(result[0]['verification_status'], 'verification_failed')

    def test_duplicate_locations_rejected(self):
        with self.assertRaises(ValueError):
            verify_docx_evidence([], [{'location_id': 'docx-p1', 'text': 'a'}, {'location_id': 'docx-p1', 'text': 'b'}])

    def test_no_permanent_document_write(self):
        with patch('builtins.open', side_effect=AssertionError('No extracted document files should be saved.')):
            data = analyze_document(UploadFile(file=BytesIO(self.contents), filename='sample.docx'))
        self.assertEqual(data['document_type'], 'docx')


class DocxBrowserTests(unittest.TestCase):
    # Reuse the same offline JS/DOM harness; do not inherit its PDF test methods.
    run_js = test_review.ReviewDashboardTests.run_js

    def setup_docx(self):
        return '''
Object.assign(upload, {filename: 'sample.docx', document_type: 'docx', ai_supported: false,
  locations: [{location_id: 'docx-p1', location_label: 'Paragraph 1', text: 'See Section 9.9.'},
    {location_id: 'docx-t1-r1-c1-p1', location_label: 'Table 1, row 1, cell 1, Paragraph 1', text: 'Use the Delivery Window.'}],
  total_locations: 2, location_notice: 'DOCX locations are not pages. AI review unavailable.',
  broken_references: [{number: '9.9', excerpt: 'See Section 9.9.', explanation: 'Missing target.',
    location_id: 'docx-p1', location_label: 'Paragraph 1'}],
  potentially_undefined_terms: [{term: 'Delivery Window', excerpt: 'Use the Delivery Window.', explanation: 'No definition detected.',
    location_id: 'docx-t1-r1-c1-p1', location_label: 'Table 1, row 1, cell 1, Paragraph 1'}]});
'''

    def test_docx_rendering_and_disabled_ai(self):
        self.run_js(self.setup_docx() + '''
await uploadFile();
assert.equal(elements.get('error').hidden, true);
assert(elements.get('summary').textContent.includes('not pages'));
assert.equal(elements.get('pages').children[0].children[0].textContent, 'Paragraph 1');
assert.equal(elements.get('ai-demo-button').disabled, true);
assert.equal(elements.get('real-button').disabled, true);
assert(!calls.includes('/ai-review-config'));
await elements.get('ai-demo-button').listeners.click();
await elements.get('real-button').listeners.click();
assert(!calls.includes('/ai-review-demo'));
assert(!calls.includes('/ai-review-real'));
assert.equal(state().length, 2);
assert(!('page' in state()[0].evidence[0]));
''', ui=True)

    def test_docx_highlight_correct_location(self):
        self.run_js(self.setup_docx() + '''
await uploadFile();
vm.runInContext('viewEvidence(reviewState.findings()[1].evidence[0])', context);
const target = elements.get('pages').children[1];
assert.equal(target.scrolled, undefined);
assert.equal(elements.get('document-scroll').scrollTop, 762.5);
assert(target.children[2].children.find(c => c.tagName === 'mark').focused);
assert(target.children[2].children.some(c => c.tagName === 'mark' && c.textContent === 'Use the Delivery Window.'));
assert(elements.get('document-navigation-status').textContent.includes('Table 1, row 1, cell 1, Paragraph 1'));
assert(!elements.get('document-navigation-status').textContent.includes('Page'));
''', ui=True)

    def test_docx_missing_and_repeated_highlights(self):
        self.run_js(self.setup_docx() + '''
upload.locations[0].text = 'See Section 9.9. See Section 9.9.';
await uploadFile();
vm.runInContext('viewEvidence(reviewState.findings()[0].evidence[0])', context);
assert(elements.get('document-navigation-status').textContent.includes('2 matches'));
assert(elements.get('document-navigation-status').textContent.includes('not identify a unique location'));
vm.runInContext('viewEvidence({locationId: "docx-p1", locationLabel: "Paragraph 1", quote: "Invented"})', context);
assert(elements.get('document-navigation-status').textContent.includes('Cannot locate'));
assert(!elements.get('pages').children[0].children[2].children.some(c => c.tagName === 'mark'));
''', ui=True)

    def test_docx_decisions_edits_and_counsel(self):
        self.run_js(self.setup_docx() + '''
await uploadFile();
vm.runInContext('reviewState.decide(reviewState.findings()[0].id, "accepted"); reviewState.edit(reviewState.findings()[0].id, "Edited target", "Please clarify the reference.")', context);
assert.equal(state()[0].decision, 'pending');
assert.equal(state()[0].evidence[0].locationId, 'docx-p1');
vm.runInContext('reviewState.decide(reviewState.findings()[0].id, "accepted"); renderReviewDashboard()', context);
elements.get('draft-counsel-button').listeners.click();
const draft = elements.get('counsel-letter').value;
assert(draft.includes('Edited target'));
assert(draft.includes('Paragraph 1'));
assert(!draft.includes('Page 1'));
assert(!draft.includes('Delivery Window'));
await elements.get('copy-counsel-button').listeners.click();
assert.equal(clipboardText, draft);
elements.get('download-counsel-button').listeners.click();
assert.equal(await blobs[0].text(), draft);
''', ui=True)

    def test_docx_failed_evidence_still_blocked_after_edit(self):
        self.run_js('''
const id = add('rules', 'verification_failed');
review.edit(id, 'Edited', 'Still failed.');
assert.equal(review.decide(id, 'accepted'), false);
''')

    def test_manipulated_docx_ai_evidence_not_in_letter(self):
        self.run_js('''
const fake = {source: 'real', category: 'ambiguity', title: 'Fake', explanation: 'Fake', decision: 'accepted', verificationStatus: 'verified',
 evidence: [{quote: 'Text', locationId: 'docx-p1', locationLabel: 'Paragraph 1', verificationStatus: 'verified'}]};
assert.equal(buildCounselLetter([fake], [{location_id: 'docx-p1', location_label: 'Paragraph 1', text: 'Text'}]).count, 0);
''')

    def test_docx_switch_to_pdf_resets_decisions_and_draft(self):
        self.run_js(self.setup_docx() + '''
await uploadFile();
vm.runInContext('reviewState.decide(reviewState.findings()[0].id, "accepted")', context);
elements.get('draft-counsel-button').listeners.click();
delete upload.document_type;
upload.broken_references = []; upload.potentially_undefined_terms = [];
await uploadFile();
assert.equal(state().length, 0);
assert.equal(elements.get('counsel-letter').value, '');
assert.equal(elements.get('ai-demo-button').disabled, false);
assert.equal(elements.get('docx-notice').hidden, true);
''', ui=True)

    def test_docx_html_injection_rendered_as_text(self):
        self.run_js(self.setup_docx() + '''
upload.locations[0].text = '<img src=x onerror=alert(1)>';
await uploadFile();
vm.runInContext('reviewState.edit(reviewState.findings()[0].id, "<script>alert(1)</script>", "<img src=x>"); renderReviewDashboard()', context);
assert.equal(elements.get('pages').children[0].children[2].textContent, '<img src=x onerror=alert(1)>');
assert(elements.get('review-findings').children[0].children[0].textContent.includes('<script>'));
assert(!html.includes('innerHTML'));
''', ui=True)

    def test_docx_file_extension_accepted_by_browser(self):
        self.run_js(self.setup_docx() + '''
elements.get('pdf-file').files = [{name: 'sample.docx', size: 100}];
await elements.get('upload-form').listeners.submit({preventDefault() {}});
assert(calls.includes('/analyze'));
''', ui=True)
